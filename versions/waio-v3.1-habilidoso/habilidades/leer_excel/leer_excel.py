import sys
import subprocess

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

import os
import argparse

def leer_excel(ruta_excel):
    if not os.path.exists(ruta_excel):
        return f"ERROR: Archivo no existe: {ruta_excel}"
    
    try:
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)
        resultado = []
        resultado.append(f"Excel leído: {ruta_excel}")
        resultado.append(f"Hojas encontradas: {len(wb.sheetnames)}")
        resultado.append("")
        
        for nombre_hoja in wb.sheetnames:
            hoja = wb[nombre_hoja]
            resultado.append(f"=== HOJA: {nombre_hoja} ===")
            
            filas = list(hoja.iter_rows(values_only=True))
            if not filas:
                resultado.append("(Hoja vacía)")
                continue
            
            # Mostrar primeras 50 filas como máximo
            for i, fila in enumerate(filas[:50], 1):
                # Convertir None a string vacío para que se vea limpio
                fila_limpia = [str(celda) if celda is not None else "" for celda in fila]
                resultado.append(f"Fila {i}: {' | '.join(fila_limpia)}")
            
            if len(filas) > 50:
                resultado.append(f"... y {len(filas) - 50} filas más")
            
            resultado.append("")
        
        return "\n".join(resultado)
    except Exception as e:
        return f"ERROR al leer Excel: {str(e)}"

def main():
    parser = argparse.ArgumentParser(description='Lee un archivo Excel y muestra su contenido')
    parser.add_argument('--ruta', required=True, help='Ruta del archivo Excel (.xlsx)')
    args = parser.parse_args()
    
    resultado = leer_excel(args.ruta)
    print(resultado)

if __name__ == "__main__":
    main()