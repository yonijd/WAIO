import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

import os
import argparse

def leer_word(ruta_word):
    if not os.path.exists(ruta_word):
        return f"ERROR: Archivo no existe: {ruta_word}"
    
    try:
        doc = docx.Document(ruta_word)
        resultado = []
        resultado.append(f"Documento Word leído: {ruta_word}")
        resultado.append("")
        
        # Extraer párrafos
        for i, parrafo in enumerate(doc.paragraphs, 1):
            if parrafo.text.strip():
                resultado.append(f"{parrafo.text}")
        
        # Extraer tablas
        if doc.tables:
            resultado.append("")
            resultado.append("--- TABLAS ---")
            for t_idx, tabla in enumerate(doc.tables, 1):
                resultado.append(f"Tabla {t_idx}:")
                for fila in tabla.rows:
                    celdas = [celda.text.strip() for celda in fila.cells]
                    resultado.append(" | ".join(celdas))
        
        if not resultado:
            return f"Documento leído: {ruta_word}\nNo se encontró texto."
        
        return "\n".join(resultado)
    except Exception as e:
        return f"ERROR al leer Word: {str(e)}"

def main():
    parser = argparse.ArgumentParser(description='Lee un archivo Word (.docx) y extrae su texto')
    parser.add_argument('--ruta', required=True, help='Ruta del archivo Word')
    args = parser.parse_args()
    
    resultado = leer_word(args.ruta)
    print(resultado)

if __name__ == "__main__":
    main()