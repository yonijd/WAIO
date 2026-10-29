import os
import sys
import argparse
import subprocess

def instalar_si_falta(modulo, paquete=None):
    if paquete is None:
        paquete = modulo
    try:
        __import__(modulo)
        return True
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", paquete])
        try:
            __import__(modulo)
            return True
        except ImportError:
            return False

def extraer_texto_archivo(ruta):
    ext = os.path.splitext(ruta)[1].lower()
    
    if ext in ['.txt', '.py', '.json', '.md', '.csv', '.xml', '.html', '.css', '.js']:
        try:
            with open(ruta, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except:
            return ""
    
    elif ext == '.pdf':
        try:
            if instalar_si_falta('PyPDF2'):
                import PyPDF2
                with open(ruta, 'rb') as f:
                    lector = PyPDF2.PdfReader(f)
                    texto = ""
                    for pagina in lector.pages:
                        texto += pagina.extract_text() + "\n"
                    return texto
        except:
            return ""
    
    elif ext == '.docx':
        try:
            if instalar_si_falta('docx', 'python-docx'):
                import docx
                doc = docx.Document(ruta)
                texto = ""
                for parrafo in doc.paragraphs:
                    texto += parrafo.text + "\n"
                return texto
        except:
            return ""
    
    elif ext in ['.xlsx', '.xls']:
        try:
            if instalar_si_falta('openpyxl'):
                import openpyxl
                wb = openpyxl.load_workbook(ruta, data_only=True)
                texto = ""
                for hoja in wb.sheetnames:
                    ws = wb[hoja]
                    texto += f"=== HOJA: {hoja} ===\n"
                    for fila in ws.iter_rows(values_only=True):
                        fila_texto = [str(c) if c is not None else "" for c in fila]
                        texto += " | ".join(fila_texto) + "\n"
                return texto
        except:
            return ""
    
    else:
        try:
            with open(ruta, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except:
            return ""

def buscar_en_archivo(ruta, patron, case_sensitive=False):
    try:
        contenido = extraer_texto_archivo(ruta)
        if not contenido:
            return []
        
        lineas = contenido.split('\n')
        resultados = []
        for i, linea in enumerate(lineas, 1):
            if case_sensitive:
                if patron in linea:
                    resultados.append((i, linea.strip()))
            else:
                if patron.lower() in linea.lower():
                    resultados.append((i, linea.strip()))
        return resultados
    except:
        return []

def buscar_en_directorio(ruta, patron, extension=None, case_sensitive=False):
    resultados = {}
    if os.path.isfile(ruta):
        encontrados = buscar_en_archivo(ruta, patron, case_sensitive)
        if encontrados:
            resultados[ruta] = encontrados
        return resultados
    
    for root, dirs, files in os.walk(ruta):
        for file in files:
            if extension and not file.endswith(extension):
                continue
            ruta_completa = os.path.join(root, file)
            encontrados = buscar_en_archivo(ruta_completa, patron, case_sensitive)
            if encontrados:
                resultados[ruta_completa] = encontrados
    return resultados

def main():
    parser = argparse.ArgumentParser(description='Busca texto en cualquier archivo (PDF, DOCX, Excel, TXT)')
    parser.add_argument('--ruta', required=True, help='Carpeta o archivo donde buscar')
    parser.add_argument('--patron', required=True, help='Texto a buscar')
    parser.add_argument('--extension', help='Filtrar por extension (ej: .pdf, .docx)')
    parser.add_argument('--case_sensitive', action='store_true', help='Sensible a mayusculas')
    
    args = parser.parse_args()
    
    resultados = buscar_en_directorio(
        args.ruta,
        args.patron,
        args.extension,
        args.case_sensitive
    )
    
    if not resultados:
        print("No se encontraron coincidencias.")
        return
    
    total = 0
    for archivo, lineas in resultados.items():
        print(f"[ARCHIVO] {archivo}")
        for num, linea in lineas:
            print(f"  linea {num}: {linea}")
            total += 1
    print(f"[TOTAL] {total} coincidencias en {len(resultados)} archivos.")

if __name__ == "__main__":
    main()