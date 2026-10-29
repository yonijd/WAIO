import sys
import subprocess

try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

import os
import argparse

def crear_word(ruta_salida, contenido, titulo=""):
    try:
        doc = Document()
        
        if titulo:
            doc.add_heading(titulo, 0)
        
        # Si contenido es string, agregar como párrafos separados por saltos de línea
        if isinstance(contenido, str):
            parrafos = contenido.split('\n')
            for parrafo in parrafos:
                if parrafo.strip():
                    doc.add_paragraph(parrafo)
        else:
            return f"ERROR: Contenido debe ser texto"
        
        doc.save(ruta_salida)
        return f"Documento Word creado: {ruta_salida}"
    except Exception as e:
        return f"ERROR al crear Word: {str(e)}"

def main():
    parser = argparse.ArgumentParser(description='Crea un archivo Word (.docx) con el texto proporcionado')
    parser.add_argument('--ruta', required=True, help='Ruta de salida del archivo Word')
    parser.add_argument('--contenido', required=True, help='Contenido del documento')
    parser.add_argument('--titulo', help='Título del documento (opcional)')
    args = parser.parse_args()
    
    resultado = crear_word(args.ruta, args.contenido, args.titulo if args.titulo else "")
    print(resultado)

if __name__ == "__main__":
    main()