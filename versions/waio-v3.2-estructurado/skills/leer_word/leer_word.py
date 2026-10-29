# Copyright (c) 2026 Yonijd
# Licensed under the MIT License.

import sys
import subprocess
import os
import argparse
import ast

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

def obtener_workdir():
    try:
        with open("temp/workdir.txt", "r") as f:
            return f.read().strip()
    except:
        return os.getcwd()

def resolver_ruta(ruta):
    workdir = obtener_workdir()
    if os.path.isabs(ruta):
        return ruta
    return os.path.join(workdir, ruta)

def leer_word(ruta_word):
    if not os.path.exists(ruta_word):
        return f"ERROR: Archivo no existe: {ruta_word}"

    try:
        doc = docx.Document(ruta_word)
        resultado = []
        resultado.append(f"Documento Word leido: {ruta_word}")
        resultado.append("")

        for i, parrafo in enumerate(doc.paragraphs, 1):
            if parrafo.text.strip():
                resultado.append(f"{parrafo.text}")

        if doc.tables:
            resultado.append("")
            resultado.append("--- TABLAS ---")
            for t_idx, tabla in enumerate(doc.tables, 1):
                resultado.append(f"Tabla {t_idx}:")
                for fila in tabla.rows:
                    celdas = [celda.text.strip() for celda in fila.cells]
                    resultado.append(" | ".join(celdas))

        if not resultado:
            return f"Documento leido: {ruta_word}\nNo se encontro texto."

        return "\n".join(resultado)
    except Exception as e:
        return f"ERROR al leer Word: {str(e)}"

def main():
    args_list = sys.argv[1:]
    if len(args_list) > 1 and args_list[0] == '--parametros':
        try:
            args_list = ast.literal_eval(args_list[1])
        except:
            args_list = []

    sys.argv = ['leer_word.py'] + args_list

    parser = argparse.ArgumentParser(description='Lee un archivo Word (.docx) y extrae su texto')
    parser.add_argument('--ruta', required=True, help='Ruta del archivo Word')
    args = parser.parse_args()

    ruta_resuelta = resolver_ruta(args.ruta)

    resultado = leer_word(ruta_resuelta)
    print(resultado)

if __name__ == "__main__":
    main()
