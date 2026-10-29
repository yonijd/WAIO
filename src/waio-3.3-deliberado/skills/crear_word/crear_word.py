# Copyright (c) 2026 Yonijd
# Licensed under the MIT License.

import sys
import subprocess
import os
import argparse
import ast

try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

def obtener_workdir():
    wd = os.environ.get("WAIO_WORKDIR")
    if not wd:
        raise RuntimeError(
            "WAIO_WORKDIR no definida. "
            "Esta skill debe ejecutarse desde el motor de WAIO."
        )
    return wd

def resolver_ruta(ruta):
    workdir = obtener_workdir()
    if os.path.isabs(ruta):
        return ruta
    return os.path.join(workdir, ruta)

def crear_word(ruta_salida, contenido, titulo=""):
    try:
        doc = Document()

        if titulo:
            doc.add_heading(titulo, 0)

        if isinstance(contenido, str):
            parrafos = contenido.split('\n')
            for parrafo in parrafos:
                if parrafo.strip():
                    doc.add_paragraph(parrafo)
        else:
            return "ERROR: Contenido debe ser texto"

        doc.save(ruta_salida)
        return f"[OK] Documento Word creado: {ruta_salida}"
    except Exception as e:
        return f"ERROR al crear Word: {str(e)}"

def main():
    args_list = sys.argv[1:]
    if len(args_list) > 1 and args_list[0] == '--parametros':
        try:
            args_list = ast.literal_eval(args_list[1])
        except:
            args_list = []

    sys.argv = ['crear_word.py'] + args_list

    parser = argparse.ArgumentParser(description='Crea un archivo Word (.docx) con el texto proporcionado')
    parser.add_argument('--ruta', required=True, help='Ruta de salida del archivo Word')
    parser.add_argument('--contenido', required=True, help='Contenido del documento')
    parser.add_argument('--titulo', help='Titulo del documento (opcional)')
    args = parser.parse_args()

    ruta_resuelta = resolver_ruta(args.ruta)

    resultado = crear_word(ruta_resuelta, args.contenido, args.titulo if args.titulo else "")
    print(resultado)
    if isinstance(resultado, str) and resultado.startswith("ERROR:"):
        sys.exit(1)
    if isinstance(resultado, str) and resultado.startswith("ERROR"):
        sys.exit(1)

if __name__ == "__main__":
    main()
