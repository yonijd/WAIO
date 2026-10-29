# Copyright (c) 2026 Yonijd
# Licensed under the MIT License.

import sys
import subprocess
import os
import argparse
import ast

try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

def obtener_workdir():
    wd = os.environ.get("WAIO_WORKDIR")
    if not wd:
        raise RuntimeError(
            "WAIO_WORKDIR no definida. "
            "Esta skill debe ejecutarse desde el motor de WAIO."
        )
    return wd

def resolver_ruta(ruta):
    workdir = obtener_workdir()
    if os.path.isabs(ruta):
        return ruta
    return os.path.join(workdir, ruta)

def leer_excel(ruta_excel):
    if not os.path.exists(ruta_excel):
        return f"ERROR: Archivo no existe: {ruta_excel}"

    try:
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)
        resultado = []
        resultado.append(f"Excel leido: {ruta_excel}")
        resultado.append(f"Hojas encontradas: {len(wb.sheetnames)}")
        resultado.append("")

        for nombre_hoja in wb.sheetnames:
            hoja = wb[nombre_hoja]
            resultado.append(f"=== HOJA: {nombre_hoja} ===")

            filas = list(hoja.iter_rows(values_only=True))
            if not filas:
                resultado.append("(Hoja vacia)")
                continue

            for i, fila in enumerate(filas[:50], 1):
                fila_limpia = [str(celda) if celda is not None else "" for celda in fila]
                resultado.append(f"Fila {i}: {' | '.join(fila_limpia)}")

            if len(filas) > 50:
                resultado.append(f"... y {len(filas) - 50} filas mas")

            resultado.append("")

        return "\n".join(resultado)
    except Exception as e:
        return f"ERROR al leer Excel: {str(e)}"

def main():
    args_list = sys.argv[1:]
    if len(args_list) > 1 and args_list[0] == '--parametros':
        try:
            args_list = ast.literal_eval(args_list[1])
        except:
            args_list = []

    sys.argv = ['leer_excel.py'] + args_list

    parser = argparse.ArgumentParser(description='Lee un archivo Excel y muestra su contenido')
    parser.add_argument('--ruta', required=True, help='Ruta del archivo Excel (.xlsx)')
    args = parser.parse_args()

    ruta_resuelta = resolver_ruta(args.ruta)

    resultado = leer_excel(ruta_resuelta)
    print(resultado)
    if isinstance(resultado, str) and resultado.startswith("ERROR"):
        sys.exit(1)

if __name__ == "__main__":
    main()
